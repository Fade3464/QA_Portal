from __future__ import annotations

from io import BytesIO
from statistics import mean
from zoneinfo import ZoneInfo
from datetime import datetime

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BODY_FONT = Font(color="111827")
STATIC_FONT = Font(color="6B7280")
GREEN_FONT = Font(color="16A34A")
PURPLE_FONT = Font(color="A855F7")
RED_FONT = Font(color="EF4444")
ORANGE_FONT = Font(color="F59E0B")
TOTAL_FILL = PatternFill("solid", fgColor="E5E7EB")
TOTAL_FONT = Font(color="111827", bold=True)
THIN_GRAY = Side(style="thin", color="D1D5DB")
BOTTOM_BORDER = Border(bottom=THIN_GRAY)

HEAT_RED = (242, 73, 92)
HEAT_YELLOW = (234, 184, 57)
HEAT_GREEN = (115, 191, 105)


TABLE_TITLES = {
    "agent": "Agent leaderboard",
    "team": "Team performance",
    "project": "Project performance",
    "client": "Client performance",
    "sold_30plus": "Sales and 30+",
    "transfers": "Transfer calls",
}


def _blend(a: tuple[int, int, int], b: tuple[int, int, int], t: float) -> str:
    t = max(0.0, min(1.0, t))
    rgb = tuple(round(x + (y - x) * t) for x, y in zip(a, b))
    return "{:02X}{:02X}{:02X}".format(*rgb)


def _heat_color(value: float, minimum: float, maximum: float) -> str:
    if maximum <= minimum:
        return _blend(HEAT_RED, HEAT_GREEN, 0.5)
    t = (value - minimum) / (maximum - minimum)
    if t <= 0.5:
        return _blend(HEAT_RED, HEAT_YELLOW, t * 2)
    return _blend(HEAT_YELLOW, HEAT_GREEN, (t - 0.5) * 2)


def _conversion_font(value: float) -> Font:
    if value >= 35:
        return GREEN_FONT
    if value >= 20:
        return ORANGE_FONT
    return RED_FONT


def _local_datetime(value: str):
    if not value:
        return ""
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return value
    if parsed.tzinfo is None:
        return parsed
    return parsed.astimezone(ZoneInfo(settings.TIME_ZONE)).replace(tzinfo=None)


def _autosize(ws, widths: dict[str, int] | None = None):
    widths = widths or {}
    for idx, cell in enumerate(ws[1], start=1):
        header = str(cell.value or "")
        if header in widths:
            ws.column_dimensions[get_column_letter(idx)].width = widths[header]
            continue
        max_len = len(header)
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            value = row[0].value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 10), 36)


def _base_sheet(title: str, headers: list[str]):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BOTTOM_BORDER
    ws.row_dimensions[1].height = 24
    return wb, ws


def _set_auto_filter(ws):
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _style_heatmap_column(ws, col_idx: int, values: list[float], *, start_row=2, end_row=None):
    if not values:
        return
    end_row = end_row or (start_row + len(values) - 1)
    minimum, maximum = min(values), max(values)
    for row_idx in range(start_row, end_row + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        try:
            numeric = float(cell.value or 0)
        except (TypeError, ValueError):
            continue
        cell.fill = PatternFill("solid", fgColor=_heat_color(numeric, minimum, maximum))
        cell.font = Font(color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")


def _append_rows(ws, rows: list[dict], headers: list[str]):
    for row in rows:
        ws.append([row.get(header, "") for header in headers])


def _agent_export(payload: dict):
    raw = list(payload.get("agent_table") or [])
    raw.sort(
        key=lambda row: (
            str(row.get("Team leader") or "").casefold(),
            -int(row.get("Sold") or 0),
            str(row.get("Agent") or "").casefold(),
        )
    )
    fixed = ["Agent", "Team leader", "Sold", "30+", "Transfers", "Conversion %"]
    dynamic = [key for key in (raw[0].keys() if raw else []) if key not in fixed]
    # Sold is hidden in the Grafana Agent Leaderboard, so omit it from XLSX too.
    headers = ["Agent", "Team leader", "30+", "Transfers", "Conversion %", *dynamic]
    wb, ws = _base_sheet(TABLE_TITLES["agent"], headers)
    _append_rows(ws, raw, headers)

    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row_idx, 3).font = PURPLE_FONT
        ws.cell(row_idx, 3).alignment = Alignment(horizontal="center")
        conv = ws.cell(row_idx, 5)
        conv.value = float(conv.value or 0) / 100
        conv.number_format = "0.00%"
        conv.font = _conversion_font(float(raw[row_idx - 2].get("Conversion %") or 0))
        conv.alignment = Alignment(horizontal="center")

    _style_heatmap_column(ws, 4, [float(row.get("Transfers") or 0) for row in raw])
    for offset, name in enumerate(dynamic, start=6):
        _style_heatmap_column(ws, offset, [float(row.get(name) or 0) for row in raw])

    # Mirror the Grafana cumulative footer as a real bold final Excel row.
    footer = [
        len({str(row.get("Agent") or "") for row in raw if str(row.get("Agent") or "")}),
        len({str(row.get("Team leader") or "") for row in raw if str(row.get("Team leader") or "")}),
        sum(int(row.get("30+") or 0) for row in raw),
        sum(int(row.get("Transfers") or 0) for row in raw),
        (mean(float(row.get("Conversion %") or 0) for row in raw) / 100) if raw else 0,
        *[sum(int(row.get(name) or 0) for row in raw) for name in dynamic],
    ]
    ws.append(footer)
    footer_row = ws.max_row
    for cell in ws[footer_row]:
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        cell.border = Border(top=THIN_GRAY)
        cell.alignment = Alignment(horizontal="center")
    ws.cell(footer_row, 5).number_format = "0.00%"

    if footer_row > 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{footer_row - 1}"
    _autosize(
        ws,
        {
            "Agent": 28,
            "Team leader": 26,
            "30+": 10,
            "Transfers": 12,
            "Conversion %": 15,
            **{name: 12 for name in dynamic},
        },
    )
    return wb


def _performance_export(payload: dict, table: str):
    if table == "team":
        rows = list(payload.get("team_table") or [])
        rows.sort(key=lambda row: (-int(row.get("Sold") or 0), str(row.get("Team") or "").casefold()))
        fixed = ["Team", "Team leader", "Sold", "30+", "Transfers", "Conversion %"]
        dynamic = [key for key in (rows[0].keys() if rows else []) if key not in fixed]
        headers = [*fixed, *dynamic]
    elif table == "project":
        source = payload.get("projects") or []
        rows = [
            {
                "Project": row.get("project", ""),
                "Sold": int(row.get("sold") or 0),
                "30+": int(row.get("thirty_plus") or 0),
                "Transfers": int(row.get("transfers") or 0),
                "Conversion %": float(row.get("conversion_pct") or 0),
            }
            for row in source
        ]
        rows.sort(key=lambda row: (-row["Sold"], str(row["Project"]).casefold()))
        headers = ["Project", "Sold", "30+", "Transfers", "Conversion %"]
        dynamic = []
    else:
        source = payload.get("presets") or []
        rows = [
            {
                "Client": row.get("preset", ""),
                "Sold": int(row.get("sold") or 0),
                "30+": int(row.get("thirty_plus") or 0),
                "Transfers": int(row.get("transfers") or 0),
                "Conversion %": float(row.get("conversion_pct") or 0),
            }
            for row in source
        ]
        rows.sort(key=lambda row: (-row["Sold"], str(row["Client"]).casefold()))
        headers = ["Client", "Sold", "30+", "Transfers", "Conversion %"]
        dynamic = []

    wb, ws = _base_sheet(TABLE_TITLES[table], headers)
    _append_rows(ws, rows, headers)

    sold_col = headers.index("Sold") + 1
    thirty_col = headers.index("30+") + 1
    transfer_col = headers.index("Transfers") + 1
    conv_col = headers.index("Conversion %") + 1
    for row_idx, row in enumerate(rows, start=2):
        ws.cell(row_idx, sold_col).font = GREEN_FONT
        ws.cell(row_idx, sold_col).alignment = Alignment(horizontal="center")
        ws.cell(row_idx, thirty_col).font = PURPLE_FONT
        ws.cell(row_idx, thirty_col).alignment = Alignment(horizontal="center")
        conv = ws.cell(row_idx, conv_col)
        raw_pct = float(row.get("Conversion %") or 0)
        conv.value = raw_pct / 100
        conv.number_format = "0.00%"
        conv.font = _conversion_font(raw_pct)
        conv.alignment = Alignment(horizontal="center")

    _style_heatmap_column(ws, transfer_col, [float(row.get("Transfers") or 0) for row in rows])
    for name in dynamic:
        col = headers.index(name) + 1
        _style_heatmap_column(ws, col, [float(row.get(name) or 0) for row in rows])

    _set_auto_filter(ws)
    _autosize(ws, {headers[0]: 28, "Team leader": 26, "Sold": 10, "30+": 10, "Transfers": 12, "Conversion %": 15, **{name: 12 for name in dynamic}})
    return wb


def _detail_export(payload: dict, table: str):
    key = "sold_details" if table == "sold_30plus" else "transfer_details"
    rows = list(payload.get(key) or [])
    if table == "sold_30plus":
        headers = ["Agent", "Team", "Number", "Client", "Sale Status", "Client status", "Disposition time"]
    else:
        headers = ["Agent", "Team", "Number", "Client", "Client status", "Disposition time"]
    wb, ws = _base_sheet(TABLE_TITLES[table], headers)

    for row in rows:
        values = [
            row.get("agent", ""),
            row.get("team", ""),
            row.get("number", ""),
            row.get("client", ""),
        ]
        if table == "sold_30plus":
            values.append(row.get("sale_status", ""))
        values.extend(
            [
                row.get("client_status", ""),
                _local_datetime(row.get("disposition_time", "")),
            ]
        )
        ws.append(values)
        dt_cell = ws.cell(ws.max_row, len(headers))
        if isinstance(dt_cell.value, datetime):
            dt_cell.number_format = "yyyy-mm-dd hh:mm:ss"

    _set_auto_filter(ws)
    _autosize(ws, {"Agent": 28, "Team": 24, "Number": 18, "Client": 20, "Sale Status": 16, "Client status": 28, "Disposition time": 22})
    return wb


def build_sales_xlsx(payload: dict, table: str) -> tuple[bytes, str]:
    table = (table or "").strip().lower()
    if table == "agent":
        wb = _agent_export(payload)
    elif table in {"team", "project", "client"}:
        wb = _performance_export(payload, table)
    elif table in {"sold_30plus", "transfers"}:
        wb = _detail_export(payload, table)
    else:
        raise ValueError("Unsupported table export")

    output = BytesIO()
    wb.save(output)
    filename = f"{table.replace('_', '-')}.xlsx"
    return output.getvalue(), filename
